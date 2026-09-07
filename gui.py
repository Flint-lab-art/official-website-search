# -*- coding: utf-8 -*-
"""Tkinter GUI：关键词导入、开始/暂停/停止/跳过、进度日志、结果表、Excel 导出"""
import os, sys, io, queue, threading, time, sqlite3
import tkinter as tk
from tkinter import ttk, filedialog, messagebox, scrolledtext

for _s in (sys.stdout, sys.stderr):
    if hasattr(_s, "reconfigure"):
        try:
            _s.reconfigure(encoding="utf-8", errors="replace")
        except Exception:
            pass

from core import config, db
from core.engine import BrowserSession
from core.baidu import run_baidu
from core.bing import run_bing

ST_LABEL = {"pending": "等待", "hit": "命中", "none": "未命中", "error": "错误"}

COLUMNS = ("kw", "bd", "bd_rank", "bd_page", "bg", "bg_rank", "bg_page", "concl")


def conclusion(bd, bg):
    b, g = bd == "hit", bg == "hit"
    if b and g:
        return "官网可见"
    if b or g:
        return "部分可见"
    if bd == "error" or bg == "error":
        return "含错误"
    return "10页未见"


class Worker(threading.Thread):
    def __init__(self, keywords, msg_q, db_path, pause_evt, stop_evt, skip_evt):
        super().__init__(daemon=True)
        self.keywords, self.msg_q = keywords, msg_q
        self.db_path = db_path
        self.pause_evt, self.stop_evt, self.skip_evt = pause_evt, stop_evt, skip_evt

    def msg(self, kind, payload):
        self.msg_q.put((kind, payload))

    def notify_captcha(self, engine, keyword):
        self.msg("captcha", (engine, keyword))

    def run(self):
        conn = sqlite3.connect(self.db_path, check_same_thread=False)
        session = BrowserSession()
        done = 0
        try:
            for kw in self.keywords:
                if self.stop_evt.is_set():
                    break
                # 暂停点
                while self.pause_evt.is_set() and not self.stop_evt.is_set():
                    time.sleep(0.3)
                if self.stop_evt.is_set():
                    break
                if self.skip_evt.is_set():
                    self.skip_evt.clear()
                self.msg("status", (kw, "百度"))
                self.msg("log", f"▶ {kw}")

                r = run_baidu(session, kw, config.SCREENSHOT_DIR,
                              skip_evt=self.skip_evt, notify=self.notify_captcha)
                db.update_result(conn, kw, "baidu", r["status"],
                                 r.get("rank"), r.get("page"),
                                 r.get("evidence"), r.get("shot"))
                self.msg("result", (kw, "baidu", r))
                self.msg("log", f"   百度 {ST_LABEL.get(r['status'], r['status'])} "
                                f"排名{r.get('rank')} 第{r.get('page')}页 | {r.get('evidence', '')[:50]}")

                if self.skip_evt.is_set():
                    # 百度侧验证码跳过 → 整个词跳过
                    self.skip_evt.clear()
                    db.update_result(conn, kw, "bing", "error", evidence="验证码跳过")
                    self.msg("result", (kw, "bing", {"status": "error", "evidence": "验证码跳过"}))
                    self.msg("log", "   已跳过（验证码）")
                else:
                    self.msg("status", (kw, "必应"))
                    g = run_bing(session, kw, config.SCREENSHOT_DIR)
                    db.update_result(conn, kw, "bing", g["status"],
                                     g.get("rank"), g.get("page"),
                                     g.get("evidence"), g.get("shot"))
                    self.msg("result", (kw, "bing", g))
                    self.msg("log", f"   必应 {ST_LABEL.get(g['status'], g['status'])} "
                                    f"排名{g.get('rank')} 第{g.get('page')}页 | {g.get('evidence', '')[:50]}")

                done += 1
                self.msg("progress", done)
        except Exception as e:
            self.msg("log", f"!! 运行异常: {e}")
        finally:
            try:
                session.close()
            except Exception:
                pass
        self.msg("done", None)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("官网检索器 · Elo / elotouch.com.cn")
        self.geometry("1180x720")
        os.makedirs(config.SCREENSHOT_DIR, exist_ok=True)

        self.msg_q = queue.Queue()
        self.pause_evt = threading.Event()
        self.stop_evt = threading.Event()
        self.skip_evt = threading.Event()
        self.worker = None
        self.conn = sqlite3.connect(config.DB_PATH)

        self._build_ui()
        self._refresh_table()
        self.after(200, self._poll)

    # ---------- UI ----------
    def _build_ui(self):
        bar = ttk.Frame(self, padding=(8, 6))
        bar.pack(fill="x")
        ttk.Button(bar, text="导入 txt", command=self._import_txt).pack(side="left")
        ttk.Button(bar, text="导入 Excel", command=self._import_excel).pack(side="left", padx=(6, 0))
        ttk.Button(bar, text="粘贴关键词", command=self._paste_keywords).pack(side="left", padx=(6, 0))
        ttk.Button(bar, text="清空任务", command=self._clear_all).pack(side="left", padx=(6, 0))
        ttk.Separator(bar, orient="vertical").pack(side="left", fill="y", padx=10)
        self.btn_start = ttk.Button(bar, text="▶ 开始", command=self._start)
        self.btn_start.pack(side="left")
        self.btn_pause = ttk.Button(bar, text="⏸ 暂停", command=self._toggle_pause, state="disabled")
        self.btn_pause.pack(side="left", padx=(6, 0))
        self.btn_skip = ttk.Button(bar, text="⏭ 跳过当前", command=self._skip, state="disabled")
        self.btn_skip.pack(side="left", padx=(6, 0))
        self.btn_stop = ttk.Button(bar, text="⏹ 停止", command=self._stop, state="disabled")
        self.btn_stop.pack(side="left", padx=(6, 0))
        ttk.Button(bar, text="导出 Excel", command=self._export).pack(side="right")

        # 状态行
        st = ttk.Frame(self, padding=(10, 2))
        st.pack(fill="x")
        self.lbl_progress = ttk.Label(st, text="进度: 0 / 0")
        self.lbl_progress.pack(side="left")
        ttk.Label(st, text="  |  ").pack(side="left")
        self.lbl_bd = ttk.Label(st, text="百度命中: 0")
        self.lbl_bd.pack(side="left")
        ttk.Label(st, text="  |  ").pack(side="left")
        self.lbl_bg = ttk.Label(st, text="必应命中: 0")
        self.lbl_bg.pack(side="left")
        ttk.Label(st, text="  |  ").pack(side="left")
        self.lbl_status = ttk.Label(st, text="就绪", foreground="#1A1B1C")
        self.lbl_status.pack(side="left")

        # 结果表
        wrap = ttk.Frame(self, padding=(8, 4))
        wrap.pack(fill="both", expand=True)
        self.tree = ttk.Treeview(wrap, columns=COLUMNS, show="headings", height=16)
        heads = {
            "kw": ("关键词", 220), "bd": ("百度", 70), "bd_rank": ("排名", 55),
            "bd_page": ("页码", 55), "bg": ("必应", 70), "bg_rank": ("排名", 55),
            "bg_page": ("页码", 55), "concl": ("结论", 90),
        }
        for c, (t, w) in heads.items():
            self.tree.heading(c, text=t)
            self.tree.column(c, width=w, anchor="w", stretch=(c == "kw"))
        vsb = ttk.Scrollbar(wrap, orient="vertical", command=self.tree.yview)
        self.tree.configure(yscrollcommand=vsb.set)
        self.tree.pack(side="left", fill="both", expand=True)
        vsb.pack(side="right", fill="y")

        # 日志
        logf = ttk.LabelFrame(self, text="日志", padding=(6, 2))
        logf.pack(fill="x", padx=8, pady=(0, 8))
        self.log = scrolledtext.ScrolledText(logf, height=9, font=("Microsoft YaHei UI", 9),
                                             state="disabled", wrap="word")
        self.log.pack(fill="x")

    def _log(self, text):
        self.log.configure(state="normal")
        self.log.insert("end", text + "\n")
        self.log.see("end")
        self.log.configure(state="disabled")

    # ---------- 数据 ----------
    def _refresh_table(self):
        self.tree.delete(*self.tree.get_children())
        rows = self.conn.execute(
            "SELECT keyword, baidu_status, baidu_rank, baidu_page, bing_status, bing_rank, bing_page "
            "FROM tasks ORDER BY id").fetchall()
        for kw, bs, br, bp, gs, gr, gp in rows:
            self.tree.insert("", "end", values=(
                kw, ST_LABEL.get(bs, bs), br if br else "", bp if bp else "",
                ST_LABEL.get(gs, gs), gr if gr else "", gp if gp else "",
                conclusion(bs, gs)))
        total = len(rows)
        hits = self.conn.execute(
            "SELECT SUM(CASE WHEN baidu_status='hit' THEN 1 ELSE 0 END), "
            "SUM(CASE WHEN bing_status='hit' THEN 1 ELSE 0 END) FROM tasks").fetchone()
        self.lbl_progress.configure(text=f"进度: 0 / {total}")
        self.lbl_bd.configure(text=f"百度命中: {hits[0] or 0}")
        self.lbl_bg.configure(text=f"必应命中: {hits[1] or 0}")

    def _update_row(self, kw, engine, r):
        for item in self.tree.get_children():
            if self.tree.item(item, "values")[0] == kw:
                vals = list(self.tree.item(item, "values"))
                if engine == "baidu":
                    vals[1] = ST_LABEL.get(r["status"], r["status"])
                    vals[2] = r.get("rank") or ""
                    vals[3] = r.get("page") or ""
                else:
                    vals[4] = ST_LABEL.get(r["status"], r["status"])
                    vals[5] = r.get("rank") or ""
                    vals[6] = r.get("page") or ""
                vals[7] = conclusion(vals[1], vals[4])
                self.tree.item(item, values=vals)
                break

    def _load_keywords(self, kws):
        kws = [k.strip() for k in kws if k.strip()]
        if not kws:
            messagebox.showinfo("提示", "没有有效关键词")
            return
        db.ensure_keywords(self.conn, kws)
        self._refresh_table()
        self._log(f"已导入 {len(kws)} 个关键词（重复自动跳过）")

    # ---------- 按钮 ----------
    def _import_txt(self):
        path = filedialog.askopenfilename(title="选择关键词 txt", filetypes=[("文本", "*.txt"), ("全部", "*.*")])
        if not path:
            return
        try:
            with open(path, encoding="utf-8") as f:
                self._load_keywords(f.readlines())
        except Exception as e:
            messagebox.showerror("错误", f"读取失败: {e}")

    def _import_excel(self):
        path = filedialog.askopenfilename(title="选择 Excel", filetypes=[("Excel", "*.xlsx *.xls")])
        if not path:
            return
        try:
            from openpyxl import load_workbook
            wb = load_workbook(path, read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            kws = [str(r[0]).strip() for r in ws.iter_rows(min_col=1, max_col=1) if r[0] is not None]
            wb.close()
            self._load_keywords(kws)
        except Exception as e:
            messagebox.showerror("错误", f"读取失败: {e}")

    def _paste_keywords(self):
        win = tk.Toplevel(self)
        win.title("粘贴关键词（每行一个）")
        win.geometry("420x420")
        txt = scrolledtext.ScrolledText(win, wrap="word")
        txt.pack(fill="both", expand=True, padx=8, pady=8)
        def ok():
            self._load_keywords(txt.get("1.0", "end").splitlines())
            win.destroy()
        ttk.Button(win, text="导入", command=ok).pack(pady=(0, 8))

    def _clear_all(self):
        if self.worker and self.worker.is_alive():
            messagebox.showwarning("提示", "运行中不能清空，先停止")
            return
        if not messagebox.askyesno("确认", "清空全部任务数据？\n（关键词、结果、进度都会删除）"):
            return
        self.conn.execute("DELETE FROM tasks")
        self.conn.commit()
        self._refresh_table()
        self._log("任务已清空")

    def _start(self):
        if self.worker and self.worker.is_alive():
            return
        pending = db.load_pending(self.conn)
        if not pending:
            messagebox.showinfo("提示", "没有待处理关键词（先导入关键词，或清空任务重跑）")
            return
        self.stop_evt.clear()
        self.pause_evt.clear()
        self.skip_evt.clear()
        self.worker = Worker(pending, self.msg_q, config.DB_PATH,
                             self.pause_evt, self.stop_evt, self.skip_evt)
        self.worker.start()
        self.btn_start.configure(state="disabled")
        self.btn_pause.configure(state="normal", text="⏸ 暂停")
        self.btn_skip.configure(state="normal")
        self.btn_stop.configure(state="normal")
        self._set_status("运行中", "#1A1B1C")
        self._log(f"开始处理 {len(pending)} 个关键词…")

    def _toggle_pause(self):
        if self.pause_evt.is_set():
            self.pause_evt.clear()
            self.btn_pause.configure(text="⏸ 暂停")
            self._set_status("继续中…", "#1A1B1C")
        else:
            self.pause_evt.set()
            self.btn_pause.configure(text="▶ 继续")
            self._set_status("已暂停（当前词完成后停止）", "#B36B00")

    def _skip(self):
        self.skip_evt.set()
        self._set_status("将跳过当前关键词…", "#B36B00")

    def _stop(self):
        self.stop_evt.set()
        self.btn_stop.configure(state="disabled")
        self._set_status("正在停止（当前词完成后退出，数据已落库）", "#B36B00")

    def _export(self):
        path = filedialog.asksaveasfilename(
            title="导出 Excel", defaultextension=".xlsx",
            filetypes=[("Excel", "*.xlsx")], initialfile="官网检索结果.xlsx")
        if not path:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
            rows = self.conn.execute(
                "SELECT keyword, baidu_status, baidu_rank, baidu_page, baidu_evidence, baidu_shot, "
                "bing_status, bing_rank, bing_page, bing_evidence, bing_shot, updated_at "
                "FROM tasks ORDER BY id").fetchall()
            wb = Workbook()
            ws = wb.active
            ws.title = "明细"
            head = ["关键词", "百度", "百度排名", "百度页码", "百度依据", "百度截图",
                    "必应", "必应排名", "必应页码", "必应依据", "必应截图", "结论", "更新时间"]
            ws.append(head)
            for c in ws[1]:
                c.font = Font(bold=True)
            for r in rows:
                kw, bs, br, bp, be, bsh, gs, gr, gp, ge, gsh, up = r
                row = [kw, ST_LABEL.get(bs, bs), br or "", bp or "", be or "", bsh or "",
                       ST_LABEL.get(gs, gs), gr or "", gp or "", ge or "", gsh or "",
                       conclusion(bs, gs), up]
                ws.append(row)
            # 截图超链接
            for idx, r in enumerate(rows, start=2):
                for col in (6, 11):
                    shot = r[col - 2]
                    if shot and os.path.exists(shot):
                        cell = ws.cell(row=idx, column=col)
                        cell.hyperlink = os.path.abspath(shot)
            for i, w in enumerate([220, 60, 55, 55, 260, 160, 60, 55, 55, 260, 160, 90, 140], start=1):
                ws.column_dimensions[chr(64 + i)].width = w

            ws2 = wb.create_sheet("汇总")
            total = len(rows)
            b_hit = sum(1 for r in rows if r[1] == "hit")
            g_hit = sum(1 for r in rows if r[6] == "hit")
            both = sum(1 for r in rows if r[1] == "hit" and r[6] == "hit")
            none = [r[0] for r in rows if r[1] == "none" and r[6] == "none"]
            err = [r[0] for r in rows if r[1] == "error" or r[6] == "error"]
            ws2.append(["指标", "数值"])
            ws2.append(["关键词总数", total])
            ws2.append(["百度官网标识命中", b_hit])
            ws2.append(["必应域名命中", g_hit])
            ws2.append(["双引擎均命中", both])
            ws2.append(["双引擎均未命中（10页内未见）", len(none)])
            ws2.append(["含错误（验证码跳过等）", len(err)])
            ws2.append([])
            ws2.append(["双引擎均未命中关键词", "、".join(none) if none else "无"])
            ws2.append(["含错误关键词", "、".join(err) if err else "无"])
            for c in ws2[1]:
                c.font = Font(bold=True)
            for i, w in enumerate([26, 80], start=1):
                ws2.column_dimensions[chr(64 + i)].width = w
            wb.save(path)
            self._log(f"已导出: {path}")
            messagebox.showinfo("完成", f"导出成功\n{path}")
        except Exception as e:
            messagebox.showerror("错误", f"导出失败: {e}")

    # ---------- 消息轮询 ----------
    def _set_status(self, text, color="#1A1B1C"):
        self.lbl_status.configure(text=text, foreground=color)

    def _poll(self):
        try:
            while True:
                kind, payload = self.msg_q.get_nowait()
                if kind == "log":
                    self._log(payload)
                elif kind == "status":
                    kw, engine = payload
                    self._set_status(f"处理中: {kw}（{engine}）")
                elif kind == "captcha":
                    engine, kw = payload
                    self._set_status(f"⚠ {engine} 要求人机验证：请在浏览器窗口完成滑块，完成后自动继续", "#C0392B")
                    self._log(f"[验证码] {engine}「{kw}」请在浏览器窗口手动验证")
                elif kind == "result":
                    kw, engine, r = payload
                    self._update_row(kw, engine, r)
                    s = self.conn.execute(
                        "SELECT SUM(CASE WHEN baidu_status='hit' THEN 1 ELSE 0 END), "
                        "SUM(CASE WHEN bing_status='hit' THEN 1 ELSE 0 END) FROM tasks").fetchone()
                    self.lbl_bd.configure(text=f"百度命中: {s[0] or 0}")
                    self.lbl_bg.configure(text=f"必应命中: {s[1] or 0}")
                elif kind == "progress":
                    n = payload
                    total = len(self.tree.get_children())
                    self.lbl_progress.configure(text=f"进度: {n} / {total}")
                elif kind == "done":
                    self._set_status("全部完成", "#1A1B1C")
                    self.btn_start.configure(state="normal")
                    self.btn_pause.configure(state="disabled")
                    self.btn_skip.configure(state="disabled")
                    self.btn_stop.configure(state="disabled")
                    s = self.conn.execute(
                        "SELECT COUNT(*), SUM(CASE WHEN baidu_status='hit' THEN 1 ELSE 0 END), "
                        "SUM(CASE WHEN bing_status='hit' THEN 1 ELSE 0 END) FROM tasks").fetchone()
                    self._log(f"===== 完成：百度命中 {s[1] or 0}/{s[0]}，必应命中 {s[2] or 0}/{s[0]} =====")
                    messagebox.showinfo("完成",
                                        f"全部处理完成\n百度官网标识命中: {s[1] or 0}/{s[0]}\n"
                                        f"必应域名命中: {s[2] or 0}/{s[0]}")
        except queue.Empty:
            pass
        self.after(200, self._poll)

    def on_close(self):
        if self.worker and self.worker.is_alive():
            self.stop_evt.set()
        self.destroy()


if __name__ == "__main__":
    # 单实例锁：防止两个 GUI 共用浏览器 profile 导致冲突
    import socket
    try:
        _lock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
        _lock.bind(("127.0.0.1", 27531))
        _lock.listen(1)
    except OSError:
        import ctypes
        ctypes.windll.user32.MessageBoxW(0, "官网检索器已在运行，请使用已打开的窗口。", "提示", 0x40)
        sys.exit(0)
    app = App()
    app.protocol("WM_DELETE_WINDOW", app.on_close)
    app.mainloop()
